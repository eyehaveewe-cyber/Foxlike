"""Append the current implementation record, preserving the original DOCX."""
from pathlib import Path
from docx import Document
from docx.shared import Inches
from docx.oxml import OxmlElement

ROOT = Path(__file__).resolve().parents[1]
TITLE = "Implementation addendum — 30 October 2026"
COMMAND_TITLE = "Companion targeting and rival coordination"
COMMAND_TEXT = "Aircraft, courier drives, walker parts, capital components and the core now share one resolved target contract for selection, orders, missiles and specialist abilities. Assignments retain blocked marks while routing around cover and regroup when the target becomes unavailable. Protect follows required friendly cargo; persistent status shows task, named target and cooldown. Guard abilities require a reachable intercept and use real flight and body collision. Healthy Shrike and Slate receive attack/cover roles; damage and withdrawal interrupt the setup without hidden player tracking or extra firing budget. The command suite adds 53 checks, including real subsystem missile hits and physical shot interception. Normal and large-text engine views were inspected."

def add_command_record(doc):
    if any(p.text == COMMAND_TITLE for p in doc.paragraphs):
        return False
    doc.add_heading(COMMAND_TITLE, 2)
    doc.add_paragraph(COMMAND_TEXT)
    return True

def main():
    path = ROOT / "docs/Foxlike_Design_and_Handoff.docx"
    doc = Document(path)
    if any(p.text == TITLE for p in doc.paragraphs):
        if add_command_record(doc):
            doc.save(path)
            print("Appended command integration record; preserved existing content")
        else:
            print("Implementation addendum and command record already exist")
        return
    doc.add_page_break()
    doc.add_heading(TITLE, 1)
    doc.add_paragraph("The version-1.0 brief and its authorship above are preserved as history. The normal game now opens a ten-operation campaign. This addendum records the implemented programming and its validation limits; it does not invent human acceptance.")
    doc.add_heading("Direction and current defaults", 2)
    doc.add_paragraph("Jonathan requested full implementation of the documentation and clarified: ‘don't worry about the human check level instructions’. Human duel/feel review therefore does not block programming. Pitch hold, reduced camera motion, restrained recoverable rival outcomes and assisted emergency behavior are reversible implementation defaults. The director has not supplied new human playtest results.")
    doc.add_heading("Playable operation order", 2)
    table = doc.add_table(rows=1, cols=2)
    table.style = "Table Grid"
    table.rows[0].cells[0].text = "Operation"
    table.rows[0].cells[1].text = "Implemented player task"
    rows = [
        ("Viridia / Blockade", "Open Warden’s crossing cannon before completing the frigate and escorting the evacuation."),
        ("Viridia / Last ferry", "Keep the same required lead ferry alive through two sectors and a sheltered transfer."),
        ("Aurora / Relay", "Use low, middle and high cover; disable RIME’s moving arms and recover the relay location."),
        ("Aurora / Courier", "Counter Shrike’s high approach; disable the courier’s rear drive and collect the codes."),
        ("Terminus / Sovereign", "Choose bay, drive, contact antenna or lance; protect extraction and escape."),
        ("Cinder / Ground corridor", "Clear three Jackal yards, choose ground or finite-fuel hover routes, disable walker parts and secure repair cargo."),
        ("Anchorage / Defense", "Protect Dragonfly and required parts, resist Rook’s bait, dismantle Harrow, clear optional recovery and depart after repair."),
        ("Terminus / Falling yard", "Arrest the runaway drive, release the coupling/stabilizer, divert the hull and leave the major wreck sections."),
        ("Tempest / Stormbreak", "Disable three corridor controls, use weather forecasts and break the finite rival flight."),
        ("Tempest / Crown", "Choose functional vanes or timed core access, escape, reassemble and fly the final bay recovery. Late retry preserves victory."),
    ]
    for label, task in rows:
        cells = table.add_row().cells
        cells[0].text, cells[1].text = label, task
        cells[0].width, cells[1].width = Inches(1.65), Inches(4.7)
        props = table.rows[-1]._tr.get_or_add_trPr()
        props.append(OxmlElement("w:cantSplit"))
    doc.add_heading("Authority and continuity", 2)
    doc.add_paragraph("operation_session.gd owns world motion, finite encounters, objectives, confirmed damage, carrier/rival/rescue facts and checkpoint boundaries. Actor, save and catalog modules keep those contracts explicit. expanded_game.gd supplies input, menus, sound, local poses and the camera. Aircraft share the Phase One vehicle controller; Jackal keeps its grounded hover controller. Presentation cannot settle an objective or grant movement.")
    doc.add_paragraph("Save version 2 maps each original six-index mission to its stable identity. New operations use new IDs. Checkpoints preserve resources, companion condition, removed systems, finite arrivals and stable emergency pose. Validation occurs before replacing live state. Atomic temporary/backup writes protect the primary; malformed and future saves offer an explicit recovery choice. Repeated result settlement is idempotent, and an earlier replay cannot replace canonical facts, frontier or the ending.")
    doc.add_paragraph("The required Cinder parts craft is used again for Anchorage delivery. Carrier hull damage and authored extraction overstrain remain separate. Recovery completion follows real craft arrival, with a holding/abort route when the lane is blocked. Crown’s defeated state is committed before the escape checkpoint. Final traffic and acknowledgments follow settled rescue, repair and rival facts.")
    doc.add_heading("Controls and presentation", 2)
    doc.add_paragraph("WASD/arrows or left stick steer; Space/LMB/RT fires. Shift/RB boosts or hovers, Ctrl/LB brakes, Q/E or A rolls. T/right-stick press selects targets, R/X launches a missile, B uses Nova, I/D-pad down interacts, and Tab/left-stick press changes formation. Tap F/Y for abilities; hold it with 1–4 or D-pad up/right/down/left for attack/protect/cover/regroup. Z/X/V or D-pad up/left/right selects the spatial stunts. Escape/Start pauses; M mutes. In-game binding labels are authoritative and button bindings are configurable.")
    doc.add_paragraph("Focus loss and controller disconnect pause safely. Inputs must return to neutral before resuming. Pitch hold, camera roll, shake and boost FOV are independent; the interface offers 100% and 150% sizing. Stunt clearance is tested before spending 30 energy. Stunts move through real world paths and do not grant defensive-roll protection. Swept ring, repair, ammunition, Nova and companion supplies retain their bounded awards and upgrades.")
    doc.add_heading("Music, models and PBR reuse", 2)
    doc.add_paragraph("Ten distinct original mission themes and six boss cues were composed using MBean1 with Sol 6.1 extra-high reasoning as requested. Editable scores, phrase source, deterministic generation, render hashes and decoded-loop validation are included. The original twelve scores and audio files remain byte-identical. Two fixed music voices crossfade over 0.65 seconds.")
    doc.add_paragraph("The ground walker is pinned to PBR blob 56c53e4afdd37745f9be8ef589150c1e9b1c71bd, with only its class name adapted before a local distance-driven pose adapter. Its root remains game-owned. A bounded damaged/abort/rejoin/alignment pose family was run in actual PBR draft-lab revision 5e34e1fc7824d3099a50deb208bffc47dc824cb0. Ninety paired checks preserve AP, cells, facing, occupancy, HP, orders and RNG with full, reduced and skipped presentation. The existing PBR lab’s 366 checks also pass. No tactical implementation or PBR merge was changed.")
    doc.add_heading("Delivery and verification limits", 2)
    doc.add_paragraph("Use Godot 4.7.2 stable and import the root project.godot, then press F5. The normal scene is scenes/main.tscn. Windows and Web exports are separate from editable source; extraction must retain the PCK beside the Windows executable, and Web files must be served over HTTP. No network account or API key is needed during play. No site publication is authorized by this work.")
    doc.add_paragraph("The native operation, physical-route and arranged journey fixtures verify objective boundaries, checkpoint reconstruction and durable/idempotent settlement. Historical compact suites preserve regression coverage. The actual browser fixture uses normal keyboard input for menu fit, music, flight, weapons, spatial maneuver and pause/resume. Dated exact results belong in docs/VERIFICATION.md; a test total is not a human completion claim.")
    doc.add_paragraph("Human feel/balance, critical listening, physical-controller operation and target Windows hardware performance remain unmeasured validation. Layout and warning values remain blockout defaults. These limitations do not reinstate the removed programming gate. Full phase-plan mapping, sources and reproduction commands are in docs/EXPANDED_CAMPAIGN.md, docs/RUNNING.md, docs/ASSET_IMPORTS.md and CREDITS.md. Library guide reads were preserved as extracted text when signed DOCX transfers were blocked by the environment’s network policy.")
    add_command_record(doc)
    doc.save(path)
    print("Updated", path)

if __name__ == "__main__": main()
